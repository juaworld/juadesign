#!/usr/bin/env python3
"""헤드리스 Chromium으로 모의 서버 기반 종단 테스트.
사용: python3 test/run_test.py  (모의 서버가 8765에서 떠 있어야 함)
"""
import io
import os
import re
import sys
import time

from PIL import Image
from playwright.sync_api import sync_playwright

BASE = os.environ.get("BASE", "http://localhost:8765")
OUT = os.environ.get("OUT", "/tmp/claude-0/sitemap-test")
os.makedirs(OUT, exist_ok=True)


def main():
    errors = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1600, "height": 1000}, locale="ko-KR")
        page.on("console", lambda m: errors.append(f"[{m.type}] {m.text}") if m.type in ("error",) else None)
        page.on("pageerror", lambda e: errors.append(f"[pageerror] {e}"))
        page.goto(f"{BASE}/?mock=1")
        page.wait_for_function("window.__app && window.__app.view")
        # 깨끗한 상태로
        page.evaluate("localStorage.clear()")
        page.reload()
        page.wait_for_function("window.__app && window.__app.view")
        time.sleep(0.8)

        # 프로젝트명·대지명
        page.fill("#projectName", "SBC리니어 충주공장 증축")
        page.press("#projectName", "Tab")
        page.fill("#siteName", "SBC 리니어")
        page.dispatch_event("#siteName", "input")

        # 검색 → 이동
        page.fill("#searchQuery", "충주시 주덕읍")
        page.click("#btnSearch")
        page.wait_for_function("document.querySelector('#log').textContent.includes('이동:')")

        # 대지 중심: 지도 클릭
        box = page.locator("#map").bounding_box()
        cx, cy = box["x"] + box["width"] / 2, box["y"] + box["height"] / 2
        page.click("#modeSite")
        page.mouse.click(cx, cy)
        page.wait_for_function("window.__app.scene.site.center !== null")

        # 필지 2개: 기존공장(중심), 증축부지(약간 동쪽)
        page.click("#modeParcel")
        page.mouse.click(cx - 10, cy - 8)
        page.wait_for_function("window.__app.scene.site.parcels.length >= 1")
        page.click("#parcelGroup .seg-btn[data-group=expansion]")
        page.mouse.click(cx + 40, cy - 8)
        page.wait_for_function("window.__app.scene.site.parcels.length >= 2")
        page.click("#modeParcel")  # 모드 종료

        # 반경이 다 보이게 → 전체 자동 생성
        page.click(".step[data-step='2'] h2")
        page.click("#btnFitRings")
        time.sleep(0.4)
        page.click(".step[data-step='3'] h2")
        page.click("#btnLoadAll")
        page.wait_for_function("document.querySelector('#log').textContent.includes('자동 생성 완료')", timeout=60000)
        time.sleep(1.0)
        page.screenshot(path=f"{OUT}/01_after_load.png")

        state = page.evaluate("""() => { const s = window.__app.scene; return {
            roads: s.roads.length, roadsVisible: s.roads.filter(r=>r.visible).length,
            roadLabels: s.roads.filter(r=>r.visible && r.label && r.label.pos).length,
            areas: s.areas.length, areasVisible: s.areas.filter(a=>a.visible).length,
            pois: s.pois.length, poisVisible: s.pois.filter(p=>p.visible).length,
            poisPlaced: s.pois.filter(p=>p.visible && p.label && p.label.pos).length,
            parcels: s.site.parcels.length,
            names: s.pois.map(p=>p.name) } }""")
        print("STATE:", state)
        log = page.inner_text("#log")
        print("LOG:\n" + log)

        # 라벨 드래그 테스트: 첫 POI 라벨을 40px 이동
        first = page.locator("#map .smap-overlay g[data-drag^='poi:']").first
        if first.count():
            pid = first.get_attribute("data-drag").split(":")[1]
            bb = first.bounding_box()
            before = page.evaluate(f"JSON.stringify(window.__app.scene.pois.find(p=>p.id==='{pid}').label.pos)")
            page.mouse.move(bb["x"] + bb["width"] / 2, bb["y"] + bb["height"] / 2)
            page.mouse.down(); page.mouse.move(bb["x"] + bb["width"] / 2 + 40, bb["y"] + bb["height"] / 2 + 10, steps=5); page.mouse.up()
            time.sleep(0.3)
            after = page.evaluate(f"JSON.stringify(window.__app.scene.pois.find(p=>p.id==='{pid}').label)")
            print("DRAG moved:", before != after, after)

        # 단계 4 열기 → 목록 확인
        page.click(".step[data-step='4'] h2")
        time.sleep(0.3)
        page.screenshot(path=f"{OUT}/02_layers.png")

        # 내보내기: PNG
        page.click(".step[data-step='6'] h2")
        page.select_option("#exportWidth", "3000")
        with page.expect_download(timeout=120000) as dl:
            page.click("#btnExportPng")
        d = dl.value
        png_path = f"{OUT}/export.png"
        d.save_as(png_path)
        im = Image.open(png_path)
        print("PNG:", im.size, os.path.getsize(png_path))
        im.thumbnail((1600, 1600)); im.save(f"{OUT}/export_thumb.png")

        # SVG (배경 포함)
        with page.expect_download(timeout=120000) as dl:
            page.click("#btnExportSvg")
        svg_path = f"{OUT}/export.svg"; dl.value.save_as(svg_path)
        svg = open(svg_path, encoding="utf-8").read()
        print("SVG bytes:", len(svg), "has image:", "<image" in svg, "texts:", svg.count("<text"))
        import xml.dom.minidom
        xml.dom.minidom.parseString(svg.encode("utf-8"))
        print("SVG well-formed: True")

        # DXF
        with page.expect_download(timeout=60000) as dl:
            page.click("#btnExportDxf")
        dxf_path = f"{OUT}/export.dxf"; dl.value.save_as(dxf_path)
        dxf = open(dxf_path, encoding="utf-8", newline="").read()
        print("DXF bytes:", len(dxf), "POLYLINE:", dxf.count("\r\nPOLYLINE"), "CIRCLE:", dxf.count("\r\nCIRCLE"), "TEXT:", dxf.count("\r\nTEXT"), "EOF:", dxf.rstrip().endswith("EOF"))
        # 그룹코드/값 쌍 검증
        lines = dxf.split("\r\n")
        if lines[-1] == "": lines = lines[:-1]
        assert len(lines) % 2 == 0, "DXF 그룹코드/값 쌍 불일치"
        for i in range(0, len(lines), 2):
            int(lines[i])

        # PPTX
        with page.expect_download(timeout=180000) as dl:
            page.click("#btnExportPptx")
        pptx_path = f"{OUT}/export.pptx"; dl.value.save_as(pptx_path)
        print("PPTX bytes:", os.path.getsize(pptx_path))
        import zipfile, xml.dom.minidom as md
        with zipfile.ZipFile(pptx_path) as z:
            names = z.namelist()
            bad = z.testzip()
            print("ZIP ok:", bad is None, "entries:", len(names))
            for n in names:
                if n.endswith(".xml") or n.endswith(".rels"):
                    md.parseString(z.read(n))
            print("all XML well-formed")
        from pptx import Presentation
        prs = Presentation(pptx_path)
        sl = prs.slides[0]
        kinds = {}
        for sh in sl.shapes:
            kinds[sh.shape_type] = kinds.get(sh.shape_type, 0) + 1
        print("PPTX slide size:", prs.slide_width, prs.slide_height, "shapes:", len(sl.shapes), {str(k): v for k, v in kinds.items()})
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text]
        print("PPTX texts sample:", texts[:8])
        sup = page.evaluate("Array.from(window.__app.suppressed||[])")
        print("SUPPRESSED:", sup)

        # 프로젝트 저장/불러오기 왕복
        with page.expect_download(timeout=60000) as dl:
            page.click("#btnSaveProject")
        js_path = f"{OUT}/project.json"; dl.value.save_as(js_path)
        page.set_input_files("#fileLoadProject", js_path)
        time.sleep(0.6)
        state2 = page.evaluate("() => ({ pois: window.__app.scene.pois.length, name: window.__app.scene.meta.name })")
        print("ROUNDTRIP:", state2)

        # 되돌리기 / 프리셋 / 인셋 줌
        n_before = page.evaluate("window.__app.scene.pois.length")
        page.evaluate("window.__app.scene.pois.pop(); window.__app.commit()")
        page.keyboard.press("Control+z")
        time.sleep(0.3)
        n_after = page.evaluate("window.__app.scene.pois.length")
        print("UNDO restored:", n_before == n_after, n_before, n_after)
        page.evaluate("window.__app.savePreset('테스트 프리셋')")
        page.evaluate("window.__app.scene.style.brightness = 0.3; window.__app.commit()")
        page.evaluate("window.__app.applyPreset('테스트 프리셋')")
        print("PRESET applied brightness:", page.evaluate("window.__app.scene.style.brightness"))
        ib = page.locator("#inset").bounding_box()
        z0 = page.evaluate("window.__app.scene.inset.zoom")
        page.mouse.move(ib["x"] + ib["width"] / 2, ib["y"] + ib["height"] / 2)
        page.mouse.wheel(0, -300)
        time.sleep(0.5)
        z1 = page.evaluate("window.__app.scene.inset.zoom")
        print("INSET wheel zoom:", z0, "->", z1)
        page.evaluate("const a = window.__app.scene.areas.find(x=>x.kind==='industrial'); if (a) { a.width = 6; window.__app.commit(); }")
        sw = page.evaluate("(() => { const p = document.querySelector('#map .smap-overlay path[stroke-dasharray]'); return p && p.getAttribute('stroke-width'); })()")
        print("AREA width applied (first dashed path stroke-width):", sw)

        # 스타일 변경 & 재배치
        page.click(".step[data-step='5'] h2")
        # 글자 굵기 (기본 7.6 → 보통 글자 + 보정 외곽선, 10 → 굵게)
        tw = page.evaluate("""() => {
          const t = Array.from(document.querySelectorAll('#map .smap-overlay text')).find(e => e.getAttribute('fill') === '#FFFFFF' && !/SITE/.test(e.textContent));
          return t && { weight: t.getAttribute('font-weight'), stroke: t.getAttribute('stroke'), sw: t.getAttribute('stroke-width'), slider: document.querySelector('#textWeight').value };
        }""")
        print("TEXT WEIGHT default:", tw)
        assert tw and tw["weight"] == "400" and tw["stroke"] == "#FFFFFF", "글자 굵기 기본값(7.6) 렌더 오류"
        page.fill("#textWeight", "10"); page.dispatch_event("#textWeight", "input")
        tw10 = page.evaluate("""() => {
          const t = Array.from(document.querySelectorAll('#map .smap-overlay text')).find(e => e.getAttribute('fill') === '#FFFFFF' && !/SITE/.test(e.textContent));
          return t && { weight: t.getAttribute('font-weight'), po: t.getAttribute('paint-order') };
        }""")
        print("TEXT WEIGHT 10:", tw10)
        assert tw10 and tw10["weight"] == "700", "글자 굵기 10 렌더 오류"
        page.fill("#textWeight", "7.6"); page.dispatch_event("#textWeight", "input")

        page.fill("#fontScale", "1.3"); page.dispatch_event("#fontScale", "input")
        page.select_option("#labelMode", "box")
        time.sleep(0.4)
        page.screenshot(path=f"{OUT}/03_style.png")

        browser.close()
    print("CONSOLE ERRORS:", len(errors))
    for e in errors[:20]:
        print("  ", e)
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
